"""Modulo per la generazione e formattazione dei commenti (popup triangolino rosso) nelle celle Excel."""

import re
from typing import Optional
from openpyxl.comments import Comment
from openpyxl.utils.cell import coordinate_to_tuple


def format_cell_coordinate(cell_ref: Optional[str]) -> str:
    """Formatta la coordinata della cella aggiungendo riga e colonna decodificate se possibile.

    Esempio:
        'E45' -> 'E45 (Riga 45, Colonna 5)'
        'FCDE!D6' -> 'FCDE!D6 (Riga 6, Colonna 4)'
        'D75+D76' -> 'D75+D76'
    """
    if not cell_ref:
        return ""
    ref = str(cell_ref).strip()
    target_coord = ref.split("!")[-1] if "!" in ref else ref
    try:
        row_idx, col_idx = coordinate_to_tuple(target_coord)
        return f"{ref} (Riga {row_idx}, Colonna {col_idx})"
    except Exception:
        return ref


def parse_source_descr(source_descr: Optional[str]) -> dict[str, Optional[str]]:
    """Estrae informazioni strutturate da una descrizione testuale sorgente preesistente."""
    result: dict[str, Optional[str]] = {
        "sheet_name": None,
        "cell_ref": None,
        "label": None,
        "detail": None,
    }
    if not source_descr:
        return result

    raw = str(source_descr).strip()
    # Se contiene pipe, analizziamo i vari pezzi
    parts = [p.strip() for p in raw.split("|") if p.strip()]

    details: list[str] = []
    for part in parts:
        lower = part.lower()
        if lower.startswith("voce "):
            result["label"] = part[5:].strip()
            continue

        if lower.startswith("foglio "):
            # es. "foglio Spese, riga 45, colonna 5 (E45)"
            # es. "foglio FCDE, cella FCDE!D6"
            # es. "foglio 1 RND a cella D57"
            after_foglio = part[7:].strip()

            # Estrai eventuale coordinata tra parentesi es. (E45)
            paren_match = re.search(r"\(([A-Za-z]+[0-9]+)\)", after_foglio)
            if paren_match:
                result["cell_ref"] = paren_match.group(1)
                sheet_part = after_foglio[: paren_match.start()]
                sheet_clean = re.sub(r",\s*riga\s+\d+.*$", "", sheet_part, flags=re.IGNORECASE).strip()
                result["sheet_name"] = sheet_clean
                continue

            # Estrai eventuale "cella XYZ"
            cella_match = re.search(r",?\s*cell[ae]\s+([^\s,]+)", after_foglio, flags=re.IGNORECASE)
            if cella_match:
                result["cell_ref"] = cella_match.group(1).strip()
                sheet_clean = after_foglio[: cella_match.start()].strip().rstrip(",")
                result["sheet_name"] = sheet_clean
                continue

            # Nessuna cella trovata, è solo il nome foglio
            result["sheet_name"] = after_foglio
            continue

        # Altre indicazioni (es. formule o note)
        details.append(part)

    if details:
        result["detail"] = " | ".join(details)

    return result


def calculate_comment_dimensions(text: Optional[str]) -> tuple[int, int]:
    """Calcola larghezza e altezza generose (in pixel) per il box del commento in Excel.

    Imposta una larghezza minima di 450px e un'altezza minima di 170px,
    adattandosi se il testo contiene linee particolarmente lunghe o numerose.
    """
    if not text:
        return 450, 170
    lines = text.split("\n")
    max_len = max(len(line) for line in lines) if lines else 30
    width = max(450, min(750, max_len * 9 + 60))
    height = max(170, len(lines) * 26 + 45)
    return width, height


def ensure_comment_dimensions(comment: Comment) -> None:
    """Imposta o ripristina dimensioni ampie e leggibili su un commento openpyxl.

    Risolve il problema del reset automatico di openpyxl a 144x79 pixel
    quando un workbook viene ricaricato da disco con load_workbook.
    """
    if comment is None:
        return
    calc_width, calc_height = calculate_comment_dimensions(comment.text)
    # Se il commento ha la dimensione di default openpyxl (144x79 o giù di lì),
    # o è comunque inferiore a quella calcolata, impostiamo la dimensione generosa
    if comment.width <= 160 or comment.height <= 90:
        comment.width = calc_width
        comment.height = calc_height
    else:
        comment.width = max(comment.width, calc_width)
        comment.height = max(comment.height, calc_height)


def fix_all_comments_in_workbook(wb) -> None:
    """Scansiona tutte le celle di tutti i fogli del workbook e corregge le dimensioni
    dei commenti (popup triangolino rosso), assicurando che non vengano salvati compressi a 144x79.
    """
    if wb is None:
        return
    for ws in wb.worksheets:
        # Usa _cells se disponibile per velocità, altrimenti iter_rows
        cells = ws._cells.values() if hasattr(ws, "_cells") and isinstance(ws._cells, dict) else (
            cell for row in ws.iter_rows() for cell in row
        )
        for cell in cells:
            if cell.comment is not None:
                ensure_comment_dimensions(cell.comment)


def create_source_comment(
    source_file_name: Optional[str] = None,
    sheet_name: Optional[str] = None,
    cell_ref: Optional[str] = None,
    label: Optional[str] = None,
    detail: Optional[str] = None,
    source_descr: Optional[str] = None,
    header: str = "FONTE DATI BDAP",
    author: str = "BDAP Automation",
) -> Comment:
    """Crea un Comment openpyxl formattato a sezioni verticali con dimensioni generose e leggibili.

    Struttura:
        📌 FONTE DATI BDAP
        • File:   nome_file.xlsx
        • Foglio: Spese
        • Cella:  E45 (Riga 45, Colonna 5)
        • Voce:   Spese per il personale
        • Note:   dettagli o formula
    """
    # Se sheet_name o cell_ref o label non sono forniti ma c'è source_descr, analizziamolo
    if source_descr:
        parsed = parse_source_descr(source_descr)
        if not sheet_name and parsed.get("sheet_name"):
            sheet_name = parsed["sheet_name"]
        if not cell_ref and parsed.get("cell_ref"):
            cell_ref = parsed["cell_ref"]
        if not label and parsed.get("label"):
            label = parsed["label"]
        if not detail and parsed.get("detail"):
            detail = parsed["detail"]
        elif detail and parsed.get("detail"):
            detail = f"{detail} | {parsed['detail']}"

    lines: list[str] = [f"📌 {header}"]

    if source_file_name:
        lines.append(f"• File:   {source_file_name}")
    if sheet_name:
        lines.append(f"• Foglio: {sheet_name}")
    if cell_ref:
        formatted_cell = format_cell_coordinate(cell_ref)
        lines.append(f"• Cella:  {formatted_cell}")
    if label:
        lines.append(f"• Voce:   {label}")
    if detail:
        lines.append(f"• Note:   {detail}")

    text = "\n".join(lines)
    comment = Comment(text, author=author)

    # Impostazione dimensioni generose
    width, height = calculate_comment_dimensions(text)
    comment.width = width
    comment.height = height

    return comment


def extract_a1_breakdown(wb_bdap) -> list[tuple[str, object]]:
    """Estrae l'elenco delle voci analitiche accantonate dal prospetto All. a1."""
    if wb_bdap is None:
        return []

    a1_sheet = None
    for name in wb_bdap.sheetnames:
        name_lower = name.lower()
        if "a1" in name_lower or "risorse acca" in name_lower or "accantonate" in name_lower:
            a1_sheet = wb_bdap[name]
            break

    if not a1_sheet:
        return []

    breakdown = []
    for r in range(12, a1_sheet.max_row + 1):
        val = a1_sheet.cell(row=r, column=9).value  # Colonna I
        if val in (None, "", 0, 0.0, "0", "0,00", "0.00"):
            continue

        col_d = str(a1_sheet.cell(row=r, column=4).value or "").strip()
        col_b = str(a1_sheet.cell(row=r, column=2).value or "").strip()

        # Ignora righe di totale complessivo
        if "totale risorse accantonate" in col_b.lower() or "totale risorse accantonate" in col_d.lower():
            continue

        if col_d:
            breakdown.append((col_d, val))
        elif col_b and not col_b.lower().startswith("totale"):
            breakdown.append((col_b, val))

    return breakdown


def create_parte_accantonata_comment(
    source_file_name: Optional[str] = None,
    sheet_name: Optional[str] = None,
    cell_ref: Optional[str] = None,
    total_formatted_val: Optional[str] = None,
    wb_bdap=None,
    header: str = "FONTE DATI BDAP",
    author: str = "BDAP Automation",
) -> Comment:
    """Crea un Comment openpyxl dettagliato per la parte accantonata del risultato di amministrazione.

    Include gli estremi della cella del totale (All. a) e, se presente nel workbook BDAP
    il prospetto All. a1, estrae e formatta analiticamente tutte le componenti non nulle,
    lasciando una riga vuota tra 'Totale accantonato' e 'Dettaglio componenti'.
    """
    lines: list[str] = [f"📌 {header}"]

    if source_file_name:
        lines.append(f"• File:   {source_file_name}")
    if sheet_name:
        lines.append(f"• Foglio: {sheet_name}")
    if cell_ref:
        formatted_cell = format_cell_coordinate(cell_ref)
        lines.append(f"• Cella:  {formatted_cell}")
    if total_formatted_val:
        tot_str = str(total_formatted_val).strip().lstrip("€").strip()
        lines.append(f"• Totale accantonato: € {tot_str}")

    breakdown = extract_a1_breakdown(wb_bdap)
    if breakdown:
        from .value_formatter import format_value_italian

        lines.append("")  # Riga vuota richiesta esplicitamente
        lines.append("• Dettaglio componenti (All. a1):")
        for item, v in breakdown:
            v_str = str(format_value_italian(v)).strip().lstrip("€").strip()
            lines.append(f"  - {item}: € {v_str}")

    text = "\n".join(lines)
    comment = Comment(text, author=author)

    width, height = calculate_comment_dimensions(text)
    comment.width = width
    comment.height = height

    return comment

