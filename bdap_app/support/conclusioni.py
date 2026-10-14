"""Modulo per la gestione e compilazione del foglio CONCLUSIONI nel template finale.

Estrae il nome del comune (per cella C2) e il numero di abitanti dell'ultimo
anno analizzato (per cella C3) dai documenti sorgente (BDAP, relazioni revisori, ecc.).
"""

from __future__ import annotations

import re
from pathlib import Path
from typing import Callable, Iterable, Optional

import openpyxl
from openpyxl.workbook import Workbook

from .comment_builder import fix_all_comments_in_workbook
from .text_utils import normalize_text

LOWERCASE_WORDS = {
    "di", "a", "da", "in", "con", "su", "per", "tra", "fra",
    "il", "lo", "la", "i", "gli", "le",
    "del", "dello", "della", "dei", "degli", "delle",
    "al", "allo", "alla", "ai", "agli", "alle",
    "nel", "nello", "nella", "nei", "negli", "nelle",
    "sul", "sullo", "sulla", "sui", "sugli", "sulle",
    "ed", "e", "od", "o",
}


def clean_comune_name(raw: str) -> str:
    """Pulisce e formatta il nome del comune rimuovendo prefissi e rumore.

    Esempio:
        'COMUNE DI GRANDOLA ED UNITI' -> 'Grandola ed Uniti'
        'COMUNE DI GARZENO' -> 'Garzeno'
        'APPIANO GENTILE (deliberazione PRSE)' -> 'Appiano Gentile'
    """
    if not raw:
        return ""

    cleaned = raw.strip()
    # Rimuovi eventuali suffissi tra parentesi es. (deliberazione PRSE) o annotazioni di demo
    cleaned = re.sub(r"\(.*?\)", "", cleaned).strip()
    cleaned = re.sub(r"[_]+demo\b", "", cleaned, flags=re.IGNORECASE).strip()
    # Rimuovi underscore usati nei percorsi cartelle
    if "_" in cleaned and " " not in cleaned:
        cleaned = cleaned.replace("_", " ").strip()

    # Rimuovi prefissi 'COMUNE DI', 'COMUNE D'', 'COMUNE DEL', ecc.
    cleaned = re.sub(r"^\s*comune\s+(?:di|d'|del|della|dei|degli|delle)?\s*", "", cleaned, flags=re.IGNORECASE).strip(" :-\t\n\r")

    if not cleaned:
        return ""

    # Se la stringa è tutta maiuscola o tutta minuscola, converti in Title Case intelligente
    if cleaned.isupper() or cleaned.islower():
        words = cleaned.split()
        title_words: list[str] = []
        for i, word in enumerate(words):
            lower_w = word.lower()
            if i > 0 and lower_w in LOWERCASE_WORDS:
                title_words.append(lower_w)
            elif "'" in word:
                parts = word.split("'")
                title_words.append("'".join(p.capitalize() for p in parts))
            else:
                title_words.append(word.capitalize())
        cleaned = " ".join(title_words)

    return cleaned


def detect_comune_from_workbook(path: Path) -> Optional[str]:
    """Estrae il nome del comune da un workbook BDAP (es. riga 7 colonna B)."""
    try:
        wb = openpyxl.load_workbook(path, data_only=True)
    except Exception:
        return None

    try:
        for ws in wb.worksheets[:3]:
            # Scansiona le prime 15 righe alla ricerca di 'denominazione ente'
            for r in range(1, min(ws.max_row or 0, 15) + 1):
                c1 = ws.cell(row=r, column=1).value
                if c1 and "denominaz" in normalize_text(str(c1)):
                    # Cerca il valore nella cella accanto o nelle prime colonne della riga
                    for col_idx in range(2, min(ws.max_column or 0, 8) + 1):
                        val = ws.cell(row=r, column=col_idx).value
                        if val and str(val).strip():
                            comune = clean_comune_name(str(val))
                            if comune:
                                return comune
    finally:
        wb.close()

    return None


def detect_comune_name(
    search_dirs: Iterable[Path],
    bdap_files: Optional[dict[int, Path]] = None,
    fallback_comune: Optional[str] = None,
) -> Optional[str]:
    """Rileva la denominazione dell'ente cercando nei file BDAP e nei workbook locali."""
    # 1. Prova prima dai file BDAP passati esplicitamente
    if bdap_files:
        for year in sorted(bdap_files.keys(), reverse=True):
            p = bdap_files[year]
            if p and p.is_file():
                detected = detect_comune_from_workbook(p)
                if detected:
                    return detected

    # 2. Cerca file Rend.*.xlsx nelle directory indicate
    for sdir in search_dirs:
        if not sdir or not sdir.exists():
            continue
        xlsx_candidates = sorted(
            list(sdir.glob("*.xlsx")) + list(sdir.glob("*/*.xlsx")),
            key=lambda f: ("rend" in f.name.lower(), f.name),
            reverse=True,
        )
        for cand in xlsx_candidates:
            if "template" in cand.name.lower():
                continue
            if "rend" in cand.name.lower() or "bdap" in str(cand).lower():
                detected = detect_comune_from_workbook(cand)
                if detected:
                    return detected

    # 3. Fallback sul parametro comune se disponibile
    if fallback_comune:
        cleaned_fallback = clean_comune_name(fallback_comune)
        if cleaned_fallback:
            return cleaned_fallback

    return None


def parse_italian_int(val_str: str) -> Optional[int]:
    """Converte una stringa numerica italiana (es. '1.300', '7.744', '660') in un intero."""
    cleaned = val_str.strip().replace(".", "").replace(",", ".").strip()
    try:
        return int(float(cleaned))
    except (ValueError, TypeError):
        return None


def _extract_text_from_pdf(pdf_path: Path) -> list[str]:
    """Estrae le pagine di testo da un PDF usando pymupdf o pypdf in fallback."""
    pages_text: list[str] = []
    try:
        import pymupdf  # type: ignore
        doc = pymupdf.open(pdf_path)
        try:
            for page in doc:
                t = page.get_text()
                if t:
                    pages_text.append(t)
        finally:
            doc.close()
        if pages_text:
            return pages_text
    except Exception:
        pass

    try:
        import pypdf  # type: ignore
        reader = pypdf.PdfReader(str(pdf_path))
        for page in reader.pages:
            t = page.extract_text()
            if t:
                pages_text.append(t)
    except Exception:
        pass

    return pages_text


def extract_population_records_from_pdf(pdf_path: Path) -> list[tuple[Optional[int], int]]:
    """Estrae tutte le coppie (anno, abitanti) trovate nel documento PDF."""
    pages_text = _extract_text_from_pdf(pdf_path)
    if not pages_text:
        return []

    file_year = None
    m_yr = re.search(r"\b(20\d\d)\b", pdf_path.name)
    if m_yr:
        file_year = int(m_yr.group(1))

    # Pattern tipici dei verbali/relazioni della Corte dei Conti e Revisori:
    pattern_year_1 = re.compile(
        r"popolazione\s+al\s+0?1[\.\-\/]0?1[\.\-\/](\d{4})[^\n\r]*?di\s+n\.?\s*([0-9\.\,]+)\s+abitanti",
        re.IGNORECASE,
    )
    pattern_year_2 = re.compile(
        r"abitanti\s+([0-9\.\,]+)[^\n\r]*?al\s+0?1[\.\-\/]0?1[\.\-\/](\d{4})",
        re.IGNORECASE,
    )
    pattern_year_3 = re.compile(
        r"comune\s+di\s+([0-9\.\,]+)\s+abitanti\s+al\s+0?1[\.\-\/]0?1[\.\-\/](\d{4})",
        re.IGNORECASE,
    )
    pattern_generic = re.compile(
        r"popolazione[^\n\r]*?di\s+n\.?\s*([0-9\.\,]+)\s+abitanti",
        re.IGNORECASE,
    )
    pattern_comune_di = re.compile(
        r"comune\s+di\s+([0-9\.\,]+)\s+abitanti",
        re.IGNORECASE,
    )
    pattern_abitanti = re.compile(
        r"abitanti\s+([0-9\.\,]+)",
        re.IGNORECASE,
    )

    records: list[tuple[Optional[int], int]] = []

    for text in pages_text:
        for m in pattern_year_1.finditer(text):
            yr = int(m.group(1))
            val = parse_italian_int(m.group(2))
            if val is not None and 10 <= val <= 5_000_000:
                records.append((yr, val))

        for m in pattern_year_2.finditer(text):
            val = parse_italian_int(m.group(1))
            yr = int(m.group(2))
            if val is not None and 10 <= val <= 5_000_000:
                records.append((yr, val))

        for m in pattern_year_3.finditer(text):
            val = parse_italian_int(m.group(1))
            yr = int(m.group(2))
            if val is not None and 10 <= val <= 5_000_000:
                records.append((yr, val))

        for m in pattern_generic.finditer(text):
            val = parse_italian_int(m.group(1))
            if val is not None and 10 <= val <= 5_000_000:
                records.append((file_year, val))

        for m in pattern_comune_di.finditer(text):
            val = parse_italian_int(m.group(1))
            if val is not None and 10 <= val <= 5_000_000:
                records.append((file_year, val))

        for m in pattern_abitanti.finditer(text):
            val = parse_italian_int(m.group(1))
            if val is not None and 10 <= val <= 5_000_000:
                records.append((file_year, val))

    return records


def extract_population_from_pdf(pdf_path: Path, target_year: Optional[int] = None) -> Optional[int]:
    """Estrae il numero di abitanti da un documento PDF (relazione revisori, pareri, ecc.)."""
    records = extract_population_records_from_pdf(pdf_path)
    if not records:
        return None

    if target_year is not None:
        target_hits = [val for yr, val in records if yr == target_year]
        if target_hits:
            return target_hits[0]

    records_with_year = [r for r in records if r[0] is not None]
    if records_with_year:
        records_with_year.sort(key=lambda r: r[0], reverse=True)
        return records_with_year[0][1]

    return records[0][1]


def detect_population(
    search_dirs: Iterable[Path],
    target_year: Optional[int] = None,
    format_with_year: bool = True,
) -> Optional[str | int]:
    """Cerca il numero di abitanti nelle cartelle del comune riportando l'ultimo anno disponibile.

    Esempio output:
        '10839 (2022)' oppure '656 (2024)'
    """
    pdf_files: list[Path] = []
    seen_paths: set[Path] = set()

    keywords = ["revisore", "parere", "istruttoria", "deliberazione", "questionario", "consuntivo", "relazione"]

    for sdir in search_dirs:
        if not sdir or not sdir.exists():
            continue
        for p in sdir.rglob("*.pdf"):
            if p not in seen_paths:
                seen_paths.add(p)
                p_lower = p.name.lower()
                if any(kw in p_lower for kw in keywords):
                    pdf_files.append(p)

    for sdir in search_dirs:
        if not sdir or not sdir.exists():
            continue
        for p in sdir.glob("*.pdf"):
            if p not in seen_paths:
                seen_paths.add(p)
                pdf_files.append(p)

    all_records: list[tuple[Optional[int], int]] = []
    for p in pdf_files:
        records = extract_population_records_from_pdf(p)
        all_records.extend(records)

    if not all_records:
        return None

    # Se c'è target_year, controlla se abbiamo un match per target_year
    if target_year is not None:
        target_hits = [r for r in all_records if r[0] == target_year]
        if target_hits:
            best_year, best_val = target_hits[0]
            if format_with_year:
                return f"{best_val} ({best_year})"
            return best_val

    # Altrimenti prendi il record con l'anno più recente disponibile
    records_with_year = [r for r in all_records if r[0] is not None]
    if records_with_year:
        records_with_year.sort(key=lambda r: r[0], reverse=True)
        best_year, best_val = records_with_year[0]
        if format_with_year:
            return f"{best_val} ({best_year})"
        return best_val

    best_val = all_records[0][1]
    return f"{best_val}" if format_with_year else best_val


def write_conclusioni_sheet(
    workbook_path: Path,
    comune_name: Optional[str] = None,
    population: Optional[int | str] = None,
    log: Optional[Callable[[str], None]] = None,
) -> bool:
    """Compila le celle C2 (comune) e C3 (abitanti) nel foglio CONCLUSIONI del template finale."""
    emit = log or (lambda _: None)

    if not workbook_path.exists():
        emit(f"File {workbook_path} non trovato per aggiornare il foglio CONCLUSIONI.")
        return False

    wb: Workbook = openpyxl.load_workbook(workbook_path)
    try:
        ws = None
        for name in wb.sheetnames:
            if name.strip().upper() == "CONCLUSIONI":
                ws = wb[name]
                break

        if ws is None:
            emit("Foglio CONCLUSIONI non trovato nel template: dati non scritti.")
            return False

        updated_items = []

        # Cella C2: Nome del comune
        if comune_name:
            c2_cell = ws["C2"]
            c2_cell.value = str(comune_name).strip()
            c2_cell.number_format = "General"
            updated_items.append(f"Comune (C2) = '{c2_cell.value}'")

        # Cella C3: Abitanti dell'ultimo anno analizzato
        c3_cell = ws["C3"]
        if population is not None:
            if isinstance(population, (int, float)):
                c3_cell.value = int(population)
                c3_cell.number_format = "#,##0"
            else:
                c3_cell.value = str(population).strip()
                c3_cell.number_format = "General"
            updated_items.append(f"Abitanti (C3) = {c3_cell.value}")
        else:
            # Se la cella conteneva il testo segnaposto del template, puliscila
            current_val = str(c3_cell.value or "").strip().lower()
            if "abitanti" in current_val or "istat" in current_val:
                c3_cell.value = None
                updated_items.append("Abitanti (C3) = [non disponibile, segnaposto rimosso]")

        fix_all_comments_in_workbook(wb)
        wb.save(workbook_path)
        if updated_items:
            emit(f"Foglio CONCLUSIONI aggiornato con successo: {', '.join(updated_items)}.")
        return True
    finally:
        wb.close()
