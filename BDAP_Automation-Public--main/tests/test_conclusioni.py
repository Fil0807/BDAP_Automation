"""Test per il modulo di supporto conclusioni (estrazione comune e abitanti)."""

from pathlib import Path
import openpyxl
import pymupdf
import pytest

from bdap_app.support.conclusioni import (
    clean_comune_name,
    parse_italian_int,
    detect_comune_from_workbook,
    extract_population_from_pdf,
    write_conclusioni_sheet,
)


def test_clean_comune_name():
    assert clean_comune_name("COMUNE DI GRANDOLA ED UNITI") == "Grandola ed Uniti"
    assert clean_comune_name("COMUNE DI GARZENO") == "Garzeno"
    assert clean_comune_name("APPIANO GENTILE (deliberazione PRSE)") == "Appiano Gentile"
    assert clean_comune_name("COMUNE DI COMO") == "Como"
    assert clean_comune_name("Grandola_Uniti_demo") == "Grandola Uniti"
    assert clean_comune_name("Garzeno") == "Garzeno"
    assert clean_comune_name("") == ""
    assert clean_comune_name("   ") == ""


def test_parse_italian_int():
    assert parse_italian_int("1.300") == 1300
    assert parse_italian_int("7.744") == 7744
    assert parse_italian_int("660") == 660
    assert parse_italian_int(" 655 ") == 655
    assert parse_italian_int("abc") is None
    assert parse_italian_int("") is None


def test_detect_comune_from_workbook(tmp_path: Path):
    wb_path = tmp_path / "Rend. 2024.xlsx"
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "1 Entrate"
    ws.cell(row=7, column=1, value="Denominazione ente:")
    ws.cell(row=7, column=2, value="COMUNE DI GRANDOLA ED UNITI")
    wb.save(wb_path)
    wb.close()

    result = detect_comune_from_workbook(wb_path)
    assert result == "Grandola ed Uniti"


def test_extract_population_from_pdf(tmp_path: Path):
    pdf_path = tmp_path / "Relazione_Revisore_2024.pdf"
    doc = pymupdf.open()
    page = doc.new_page()
    page.insert_text(
        (50, 100),
        "1.1. Verifiche preliminari\n"
        "L'Ente registra una popolazione al 01.01.2024, ai sensi dell'art.156, comma 2, del TUEL, di n.\n"
        "1.300 abitanti.\n"
        "L'Ente non è in dissesto;"
    )
    doc.save(pdf_path)
    doc.close()

    pop_2024 = extract_population_from_pdf(pdf_path, target_year=2024)
    assert pop_2024 == 1300

    pop_any = extract_population_from_pdf(pdf_path)
    assert pop_any == 1300


def test_write_conclusioni_sheet(tmp_path: Path):
    wb_path = tmp_path / "Template_Analisi.xlsx"
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "CONCLUSIONI"
    ws["B2"] = "COMUNE DI:"
    ws["C2"] = "Nome"
    ws["B3"] = "ABITANTI:"
    ws["C3"] = "numero abitanti Istat 01.01.2026"
    wb.save(wb_path)
    wb.close()

    success = write_conclusioni_sheet(
        workbook_path=wb_path,
        comune_name="Grandola ed Uniti",
        population=1300,
    )
    assert success is True

    # Verifica i valori scritti nel file
    wb_read = openpyxl.load_workbook(wb_path)
    ws_read = wb_read["CONCLUSIONI"]
    assert ws_read["C2"].value == "Grandola ed Uniti"
    assert ws_read["C2"].number_format == "General"
    assert ws_read["C3"].value == 1300
    assert ws_read["C3"].number_format == "#,##0"
    wb_read.close()


def test_write_conclusioni_sheet_no_population(tmp_path: Path):
    wb_path = tmp_path / "Template_Analisi.xlsx"
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "CONCLUSIONI"
    ws["B2"] = "COMUNE DI:"
    ws["C2"] = "Nome"
    ws["B3"] = "ABITANTI:"
    ws["C3"] = "numero abitanti Istat 01.01.2026"
    wb.save(wb_path)
    wb.close()

    success = write_conclusioni_sheet(
        workbook_path=wb_path,
        comune_name="Garzeno",
        population=None,
    )
    assert success is True

    wb_read = openpyxl.load_workbook(wb_path)
    ws_read = wb_read["CONCLUSIONI"]
    assert ws_read["C2"].value == "Garzeno"
    # Il testo dummy deve essere stato pulito
    assert ws_read["C3"].value is None
    wb_read.close()


def test_detect_comune_and_population_integration(tmp_path: Path):
    from bdap_app.support.conclusioni import detect_comune_name, detect_population

    # Setup directory structure: comune_root/1 doc/dati BDAP
    comune_dir = tmp_path / "Comune_Test_demo"
    bdap_dir = comune_dir / "1 doc" / "dati BDAP"
    bdap_dir.mkdir(parents=True)

    # 1. BDAP file
    bdap_file = bdap_dir / "Rend. 2024.xlsx"
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.cell(row=7, column=1, value="Denominazione ente:")
    ws.cell(row=7, column=2, value="COMUNE DI TEST INTEGRATO")
    wb.save(bdap_file)
    wb.close()

    # 2. PDF file with population
    pdf_file = comune_dir / "1 doc" / "Relazione Revisore 2024.pdf"
    doc = pymupdf.open()
    page = doc.new_page()
    page.insert_text((50, 100), "L'Ente registra una popolazione al 01.01.2024, ai sensi dell'art.156, comma 2, del TUEL, di n. 5.432 abitanti.")
    doc.save(pdf_file)
    doc.close()

    # Test detection
    comune_name = detect_comune_name(
        search_dirs=[comune_dir, bdap_dir],
        bdap_files={2024: bdap_file},
        fallback_comune="Comune_Test_demo",
    )
    assert comune_name == "Test Integrato"

    pop = detect_population(
        search_dirs=[comune_dir, bdap_dir],
        target_year=2024,
    )
    assert pop == 5432

